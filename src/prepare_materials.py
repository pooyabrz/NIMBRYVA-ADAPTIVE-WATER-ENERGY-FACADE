"""Reproduce ETH sheet extraction; optional dependency openpyxl, no browser.

Download bytes must match the frozen source hash. Does not overwrite active
model parameters. Preserve the CC BY-NC dataset attribution in NOTICE.md.
"""
import argparse,csv,hashlib,json,urllib.request
from pathlib import Path
from assembly import ROOT


def main():
    p=argparse.ArgumentParser();p.add_argument('--xlsx',type=Path,help='Existing author dataset; otherwise download official exact bitstream');a=p.parse_args()
    import openpyxl
    meta=json.loads((ROOT/'config/evidence/new-source-provenance.json').read_text())['ETH_FIGURES']
    if a.xlsx:raw=a.xlsx.read_bytes()
    else:
        with urllib.request.urlopen(meta['url'],timeout=60) as r:raw=r.read()
    if hashlib.sha256(raw).hexdigest()!=meta['sha256_original_xlsx']:
        raise ValueError('Source workbook changed; inspect/version evidence before replacing extracted data')
    import io
    w=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True);out=ROOT/'data/materials'
    with (out/'haechler-fig1c-raw.csv').open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['wavelength_um','absorptance_raw','physical_0_1','trusted_0p25_20um'])
        for row in w['Figure 1C'].iter_rows(min_row=2,values_only=True):
            if isinstance(row[0],(float,int)) and isinstance(row[1],(float,int)):
                lam,alpha=row[0],row[1]/100;writer.writerow([lam,alpha,int(0<=alpha<=1),int(.25<=lam<=20 and 0<=alpha<=1)])
    for sheet,name,cols in [('Figure 4B','haechler-fig4b.csv',10),('Figure 4C','haechler-fig4c.csv',3)]:
        with (out/name).open('w',newline='') as f:
            writer=csv.writer(f)
            for row in w[sheet].iter_rows(values_only=True):writer.writerow(row[:cols])
    print('Reproduced frozen author-data extraction; read data/materials/NOTICE.md.')


if __name__=='__main__':main()
